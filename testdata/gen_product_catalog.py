from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

out = Path(__file__).with_name("产品目录-测试50.xlsx")

categories = [
    ("注塑模具", "JM"),
    ("冲压模具", "CY"),
    ("夹具治具", "JJ"),
    ("精密零件", "JL"),
    ("包装辅材", "BZ"),
]

# name, category_idx, model_suffix, material, unit, price, moq, lead_days, status, note
products = [
    ("手机外壳注塑模", 0, "A01", "P20钢", "套", 128000, 1, 35, "在售", "含热流道"),
    ("家电面板注塑模", 0, "A02", "H13钢", "套", 96000, 1, 30, "在售", ""),
    ("汽车内饰件注塑模", 0, "A03", "718钢", "套", 185000, 1, 45, "在售", "双色模可选"),
    ("瓶盖注塑模", 0, "A04", "S136钢", "套", 52000, 1, 25, "在售", "24腔"),
    ("医疗耗材注塑模", 0, "A05", "S136H", "套", 168000, 1, 40, "在售", "洁净要求"),
    ("连接器注塑模", 0, "A06", "NAK80", "套", 78000, 1, 28, "在售", ""),
    ("日用品收纳盒模", 0, "A07", "P20钢", "套", 45000, 1, 22, "在售", ""),
    ("玩具壳体注塑模", 0, "A08", "718钢", "套", 38000, 1, 20, "在售", ""),
    ("灯具灯罩注塑模", 0, "A09", "S136钢", "套", 62000, 1, 26, "在售", "透光件"),
    ("电子按键注塑模", 0, "A10", "NAK80", "套", 41000, 1, 18, "在售", ""),
    ("金属拉伸冲压模", 1, "B01", "Cr12MoV", "套", 72000, 1, 30, "在售", ""),
    ("五金连续冲压模", 1, "B02", "SKD11", "套", 98000, 1, 35, "在售", "级进模"),
    ("端子冲切模", 1, "B03", "SKH51", "套", 56000, 1, 25, "在售", ""),
    ("钣金折弯模", 1, "B04", "Cr12", "套", 34000, 1, 20, "在售", ""),
    ("铝材冲孔模", 1, "B05", "DC53", "套", 48000, 1, 22, "在售", ""),
    ("不锈钢落料模", 1, "B06", "SKD11", "套", 67000, 1, 28, "在售", ""),
    ("弹簧片成形模", 1, "B07", "SKH9", "套", 59000, 1, 26, "在售", ""),
    ("连接片级进模", 1, "B08", "ASP23", "套", 112000, 1, 40, "在售", ""),
    ("盖板冲裁模", 1, "B09", "Cr12MoV", "套", 39000, 1, 18, "在售", ""),
    ("散热片冲压模", 1, "B10", "DC53", "套", 74000, 1, 30, "在售", ""),
    ("CNC定位夹具", 2, "C01", "铝合金6061", "套", 8600, 1, 12, "在售", ""),
    ("焊接工装夹具", 2, "C02", "Q235+定位销", "套", 15200, 1, 15, "在售", ""),
    ("检具（通止规）", 2, "C03", "合金钢", "套", 3200, 1, 10, "在售", ""),
    ("气动夹具总成", 2, "C04", "铝合金+气缸", "套", 12800, 1, 14, "在售", ""),
    ("装配定位治具", 2, "C05", "POM+铝", "套", 5400, 1, 8, "在售", ""),
    ("激光刻字治具", 2, "C06", "铝合金", "套", 4100, 1, 7, "在售", ""),
    ("视觉检测治具", 2, "C07", "铝合金阳极", "套", 9800, 1, 12, "在售", ""),
    ("螺丝锁付治具", 2, "C08", "钢+铝", "套", 6700, 1, 10, "在售", ""),
    ("测试探针治具", 2, "C09", "Peek+铜针", "套", 11500, 1, 14, "在售", ""),
    ("周转托盘治具", 2, "C10", "ABS/PC", "套", 2800, 5, 7, "在售", ""),
    ("精密销轴", 3, "D01", "SUJ2", "件", 12.5, 100, 5, "在售", "淬火研磨"),
    ("导向套", 3, "D02", "铜合金", "件", 18.0, 50, 5, "在售", ""),
    ("顶针组件", 3, "D03", "SKH51", "件", 6.8, 200, 4, "在售", ""),
    ("滑块耐磨块", 3, "D04", "青铜", "件", 45.0, 20, 6, "在售", ""),
    ("定位销", 3, "D05", "SKD61", "件", 8.2, 100, 4, "在售", ""),
    ("弹簧（模具用）", 3, "D06", "SWOSC-V", "件", 3.5, 500, 3, "在售", ""),
    ("水嘴接头", 3, "D07", "黄铜", "件", 4.2, 200, 3, "在售", ""),
    ("热流道喷嘴", 3, "D08", "H13", "件", 680, 2, 10, "在售", ""),
    ("模架标准件", 3, "D09", "S50C", "套", 2200, 1, 8, "在售", "A型"),
    ("密封圈套装", 3, "D10", "氟橡胶", "套", 56, 10, 3, "在售", ""),
    ("防静电周转箱", 4, "E01", "导电PP", "个", 28, 50, 5, "在售", ""),
    ("珍珠棉内托", 4, "E02", "EPE", "套", 3.6, 1000, 7, "在售", "定制开模"),
    ("纸箱（出口）", 4, "E03", "五层瓦楞", "个", 4.8, 500, 5, "在售", ""),
    ("气泡膜卷材", 4, "E04", "PE", "卷", 65, 20, 3, "在售", "宽1.2m"),
    ("标签贴纸", 4, "E05", "PET", "卷", 120, 10, 4, "在售", "可打印"),
    ("干燥剂（工业）", 4, "E06", "硅胶", "箱", 85, 5, 3, "在售", "500包/箱"),
    ("缠绕膜", 4, "E07", "LLDPE", "卷", 42, 30, 3, "在售", ""),
    ("纸护角", 4, "E08", "牛皮纸", "根", 1.2, 1000, 4, "在售", ""),
    ("吸塑托盘", 4, "E09", "PET", "个", 2.8, 2000, 10, "停产", "清仓"),
    ("缓冲泡沫块", 4, "E10", "EPS", "块", 0.9, 5000, 5, "在售", ""),
]

assert len(products) == 50

headers = [
    "序号",
    "产品编号",
    "产品名称",
    "产品分类",
    "规格型号",
    "材质/参数",
    "单位",
    "参考单价(元)",
    "最小起订量",
    "交期(天)",
    "状态",
    "备注",
]

wb = Workbook()
ws = wb.active
ws.title = "产品目录"

header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_font = Font(name="微软雅黑", size=10)
cell_align = Alignment(horizontal="left", vertical="center")
center_align = Alignment(horizontal="center", vertical="center")
thin = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
alt_fill = PatternFill("solid", fgColor="F2F7FB")

ws.append(headers)
ws.row_dimensions[1].height = 28
for col in range(1, len(headers) + 1):
    cell = ws.cell(1, col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin

for i, (name, cat_idx, model, material, unit, price, moq, lead, status, note) in enumerate(
    products, 1
):
    cat_name, prefix = categories[cat_idx]
    code = f"KY-{prefix}-{i:03d}"
    model_full = f"{prefix}-{model}"
    ws.append(
        [i, code, name, cat_name, model_full, material, unit, price, moq, lead, status, note]
    )
    for col in range(1, len(headers) + 1):
        cell = ws.cell(i + 1, col)
        cell.font = cell_font
        cell.border = thin
        cell.alignment = center_align if col in (1, 2, 4, 5, 7, 8, 9, 10, 11) else cell_align
        if i % 2 == 0:
            cell.fill = alt_fill
        if col == 8:
            cell.number_format = "#,##0.00" if isinstance(price, float) else "#,##0"
    ws.row_dimensions[i + 1].height = 20

widths = [8, 16, 22, 14, 14, 16, 8, 14, 12, 10, 8, 16]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.auto_filter.ref = f"A1:L{len(products) + 1}"
ws.freeze_panes = "A2"

ws2 = wb.create_sheet("分类汇总")
ws2.append(["产品分类", "数量", "在售", "停产"])
for col in range(1, 5):
    c = ws2.cell(1, col)
    c.font = header_font
    c.fill = header_fill
    c.alignment = header_align

by_cat: dict[str, dict[str, int]] = defaultdict(lambda: {"total": 0, "在售": 0, "停产": 0})
for item in products:
    cat_name = categories[item[1]][0]
    status = item[8]
    by_cat[cat_name]["total"] += 1
    by_cat[cat_name][status] += 1

for r, (cat, d) in enumerate(by_cat.items(), 2):
    ws2.cell(r, 1, cat)
    ws2.cell(r, 2, d["total"])
    ws2.cell(r, 3, d["在售"])
    ws2.cell(r, 4, d["停产"])

for i, w in enumerate([14, 10, 10, 10], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

wb.save(out)
print(f"written: {out}")
print(f"rows: {len(products)}")
